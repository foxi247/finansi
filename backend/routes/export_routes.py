from fastapi import APIRouter, Depends, Header, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from database import get_db
from models import User, Transaction
from routes.transactions import get_user
from datetime import datetime
from typing import Optional
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

router = APIRouter()


@router.get("/excel")
def export_excel(
    date_from: Optional[datetime] = None,
    date_to: Optional[datetime] = None,
    user: User = Depends(get_user),
    db: Session = Depends(get_db)
):
    q = db.query(Transaction).filter(Transaction.user_id == user.id)
    if date_from:
        q = q.filter(Transaction.date >= date_from)
    if date_to:
        q = q.filter(Transaction.date <= date_to)

    transactions = q.order_by(Transaction.date.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Транзакции"

    # Стили
    header_fill = PatternFill(start_color="1E1E35", end_color="1E1E35", fill_type="solid")
    income_fill = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
    expense_fill = PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC')
    )

    # Заголовок
    ws.merge_cells("A1:F1")
    title_cell = ws["A1"]
    title_cell.value = f"Finansi — Отчёт по транзакциям"
    title_cell.font = Font(bold=True, size=14, color="1E1E35")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    # Период
    ws.merge_cells("A2:F2")
    period_cell = ws["A2"]
    if date_from or date_to:
        period_cell.value = f"Период: {date_from.strftime('%d.%m.%Y') if date_from else '—'} — {date_to.strftime('%d.%m.%Y') if date_to else '—'}"
    else:
        period_cell.value = f"Все транзакции | Экспорт: {datetime.now().strftime('%d.%m.%Y %H:%M')}"
    period_cell.font = Font(size=10, color="666666")
    period_cell.alignment = Alignment(horizontal="center")

    # Пустая строка
    ws.row_dimensions[3].height = 5

    # Заголовки столбцов
    headers = ["#", "Дата", "Тип", "Категория", "Сумма (₽)", "Заметка"]
    col_widths = [5, 15, 12, 20, 15, 30]

    for col_idx, (header, width) in enumerate(zip(headers, col_widths), 1):
        cell = ws.cell(row=4, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[4].height = 25

    # Данные
    total_income = 0
    total_expense = 0

    for row_idx, tx in enumerate(transactions, 5):
        is_income = tx.type == "income"
        row_fill = income_fill if is_income else expense_fill
        sign = "+" if is_income else "-"

        values = [
            row_idx - 4,
            tx.date.strftime("%d.%m.%Y %H:%M"),
            "Доход" if is_income else "Расход",
            tx.category.name if tx.category else "—",
            tx.amount,
            tx.note or "—"
        ]

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.fill = row_fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=(col_idx == 6))
            if col_idx == 5:
                cell.number_format = '#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")

        if is_income:
            total_income += tx.amount
        else:
            total_expense += tx.amount

    # Итоговые строки
    summary_row = len(transactions) + 5 + 1
    ws.merge_cells(f"A{summary_row}:D{summary_row}")
    ws[f"A{summary_row}"].value = "ИТОГО ДОХОДЫ:"
    ws[f"A{summary_row}"].font = Font(bold=True, color="155724")
    ws[f"E{summary_row}"].value = total_income
    ws[f"E{summary_row}"].font = Font(bold=True, color="155724")
    ws[f"E{summary_row}"].number_format = '#,##0.00'

    summary_row += 1
    ws.merge_cells(f"A{summary_row}:D{summary_row}")
    ws[f"A{summary_row}"].value = "ИТОГО РАСХОДЫ:"
    ws[f"A{summary_row}"].font = Font(bold=True, color="721C24")
    ws[f"E{summary_row}"].value = total_expense
    ws[f"E{summary_row}"].font = Font(bold=True, color="721C24")
    ws[f"E{summary_row}"].number_format = '#,##0.00'

    summary_row += 1
    ws.merge_cells(f"A{summary_row}:D{summary_row}")
    ws[f"A{summary_row}"].value = "БАЛАНС:"
    ws[f"A{summary_row}"].font = Font(bold=True, size=12)
    ws[f"E{summary_row}"].value = total_income - total_expense
    balance_color = "155724" if total_income >= total_expense else "721C24"
    ws[f"E{summary_row}"].font = Font(bold=True, size=12, color=balance_color)
    ws[f"E{summary_row}"].number_format = '#,##0.00'

    # Freeze header
    ws.freeze_panes = "A5"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"finansi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
